"""stage 2 tests: docs and prose inputs, url sources, doc-site crawls.
no network: the crawler and firecrawl run on injected transports."""

import base64
import json

import httpx
import pytest

from app.ingest.docsite import crawl
from app.ingest.docs import parse_document, parse_direct_text, suffix_of
from app.ingest.firecrawl import FirecrawlClient
from app.ingest.html import extract_links, html_to_text, same_root


# ---------- html ----------

def test_html_to_text_drops_scripts_and_decodes_entities():
    html = "<html><script>evil()</script><body><h1>fast install</h1><p>safe &amp; reproducible</p></body></html>"
    text = html_to_text(html)
    assert "evil" not in text
    assert "fast install" in text and "safe & reproducible" in text


def test_extract_links_resolves_and_defragments():
    html = '<a href="/guide">a</a> <a href="https://x.dev/#top">b</a> <a href="mailto:x@y.z">c</a> <a href="javascript:void(0)">d</a> <a href="/guide">dup</a>'
    links = extract_links(html, "https://x.dev/docs/")
    assert "https://x.dev/guide" in links  # root-relative resolves against the host
    assert "https://x.dev/" in links  # fragment stripped, root stays root
    assert all("#" not in l for l in links)
    assert all(not l.startswith("mailto:") and not l.startswith("javascript:") for l in links)
    assert len(links) == len(set(links))


def test_same_root_boundary():
    assert same_root("https://x.dev/docs/", "https://x.dev/docs/page")
    assert not same_root("https://x.dev/docs/", "https://x.dev/blog")  # off root
    assert not same_root("https://x.dev/docs/", "https://other.dev/docs/")  # off host


# ---------- doc-site crawl ----------

def _site_transport() -> httpx.MockTransport:
    pages = {
        "https://x.dev/docs/": '<html><a href="/docs/intro">intro</a><a href="/docs/deep">deep</a><a href="https://ads.example.com/track">ad</a></html>',
        "https://x.dev/docs/intro": '<html><a href="/docs/deep#part">deep again</a><a href="/blog">blog</a></html>',
        "https://x.dev/docs/deep": "<html><p>the engine rebuilds geometry deterministically</p></html>",
    }

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if url in pages:
            return httpx.Response(200, text=pages[url], headers={"content-type": "text/html"})
        return httpx.Response(404)

    return httpx.MockTransport(handler)


async def test_crawl_stays_under_the_root_and_budget():
    async with httpx.AsyncClient(transport=_site_transport()) as client:
        result = await crawl("https://x.dev/docs/", budget=10, client=client)
    urls = [p["url"] for p in result["pages"]]
    assert urls == ["https://x.dev/docs/", "https://x.dev/docs/intro", "https://x.dev/docs/deep"]
    assert "https://ads.example.com/track" not in urls
    report = result["report"]
    assert report["visited"] == 3 and report["skipped_external"] >= 1
    assert report["skipped_off_root"] >= 1  # /blog refused: outside the doc root
    assert not report["unread_budget"]


def test_extract_links_skips_assets():
    """the crawler walks pages, not css and js. assets burn budget and
    pollute the failure report with non-pages."""
    from app.ingest.html import extract_links

    html = '<a href="/docs">docs</a><link href="/styles.css"/><script src="/app.js"></script>' \
           '<a href="/logo.svg"/><a href="/docs/intro?v=1">intro</a>'
    assert extract_links(html, "https://site.dev") == ["https://site.dev/docs", "https://site.dev/docs/intro?v=1"]


async def test_crawl_reports_fetch_failures():
    """a dead host must say so: visited 0 with no reason is a mystery,
    visited 0 with fetch_failures is a map."""

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(503)

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        result = await crawl("https://dead.example", budget=3, client=client)
    assert result["pages"] == []
    assert result["report"]["visited"] == 0
    assert result["report"]["fetch_failures"] == ["https://dead.example: status 503"]


async def test_crawl_dedups_root_and_its_slash():
    """https://site and https://site/ are one page, not two sources."""

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, text='<a href="/">home</a><a href="/docs">docs</a>', headers={"content-type": "text/html"})

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        result = await crawl("https://site.dev/", budget=4, client=client)
    urls = [p["url"] for p in result["pages"]]
    assert urls.count("https://site.dev") == 1
    assert "https://site.dev/docs" in urls


async def test_crawl_budget_reports_unread():
    async with httpx.AsyncClient(transport=_site_transport()) as client:
        result = await crawl("https://x.dev/docs/", budget=2, client=client)
    assert len(result["pages"]) == 2
    assert result["report"]["unread_budget"]  # a map that names what it did not read


# ---------- docs ----------

def test_direct_parse_chunks_with_doc_provenance():
    chunks = parse_direct_text("# the project\nrenders the exact product mesh", "readme.md")
    assert chunks[0]["kind"] == "doc" and chunks[0]["path"] == "readme.md"


def test_suffix_of():
    assert suffix_of("Spec.PDF") == ".pdf"
    assert suffix_of("noext") == ""


async def test_parse_document_rejects_misuse():
    with pytest.raises(ValueError, match="as text"):
        await parse_document("notes.md", content_b64="aGk=")
    with pytest.raises(ValueError, match="base64"):
        await parse_document("bad.pdf", content_b64="not-base64!!")


def _firecrawl() -> FirecrawlClient:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"data": {"markdown": "# parsed pdf\nthe project ships a deterministic engine"}})
    return FirecrawlClient(api_key="test-key", transport=httpx.MockTransport(handler))


def _one_page_pdf(text: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 8, text)
    return pdf.output()


async def test_pdf_parses_locally_with_page_provenance():
    """no external service: pypdf runs where beholder runs, and the page
    number survives into the chunk."""
    raw = _one_page_pdf("the project ships a deterministic engine. " * 30)
    chunks = await parse_document("spec.pdf", content_b64=base64.b64encode(raw).decode())
    assert chunks and chunks[0]["kind"] == "doc"
    assert chunks[0]["page"] == 1
    assert any("deterministic engine" in c["text"] for c in chunks)


async def test_two_page_pdf_cites_both_pages():
    raw = _one_page_pdf("page one text. " + "filler line. " * 40)
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    r1 = PdfReader(__import__("io").BytesIO(raw))
    for p in r1.pages:
        writer.add_page(p)
    writer.add_page(r1.pages[0])
    buf = __import__("io").BytesIO()
    writer.write(buf)
    chunks = await parse_document("spec.pdf", content_b64=base64.b64encode(buf.getvalue()).decode())
    assert {c["page"] for c in chunks} == {1, 2}


async def test_docx_parses_locally():
    import docx as docxlib

    doc = docxlib.Document()
    doc.add_paragraph("the studio renders product photos deterministically")
    doc.add_paragraph("no diffusion model touches the product")
    buf = __import__("io").BytesIO()
    doc.save(buf)
    chunks = await parse_document("brief.docx", content_b64=base64.b64encode(buf.getvalue()).decode())
    assert chunks and chunks[0]["kind"] == "doc"
    assert "deterministically" in chunks[0]["text"]


async def test_broken_pdf_is_a_clean_rejection():
    with pytest.raises(ValueError):
        await parse_document("spec.pdf", content_b64=base64.b64encode(b"%PDF-not-a-real-file").decode())


def test_firecrawl_without_key_is_honest():
    client = FirecrawlClient(api_key=None)
    import asyncio

    async def go():
        try:
            await client.scrape_markdown("https://x.dev/paper.pdf")
            return False
        except RuntimeError as e:
            return "BEHOLDER_FIRECRAWL_KEY" in str(e)

    assert asyncio.run(go())


def test_semantic_selection_prioritizes_prose_and_round_robins():
    """text beats doc beats url beats repo; sources of equal priority
    alternate, so one long doc cannot eat the whole budget."""
    from app.brain.extraction import select_semantic_chunks

    def chunk(source_id, source_kind, idx):
        return {"source_id": source_id, "source_kind": source_kind, "path": f"{source_kind}-{idx}", "text": "x"}

    rows = (
        [chunk("r1", "repo", i) for i in range(6)]
        + [chunk("u1", "text", 0)]
        + [chunk("d2", "doc", i) for i in range(5)]
        + [chunk("d1", "doc", i) for i in range(5)]
    )
    picked = select_semantic_chunks(rows, 6)
    kinds = [(r["source_kind"], r["source_id"]) for r in picked]
    assert kinds == [("text", "u1"), ("doc", "d1"), ("doc", "d2"), ("doc", "d1"), ("doc", "d2"), ("doc", "d1")]
    # the repo source is never read while prose sources still have chunks
    assert all(r["source_kind"] != "repo" for r in picked)
    # within a source, stored order is kept
    d1 = [r["path"] for r in picked if r["source_id"] == "d1"]
    assert d1 == ["doc-0", "doc-1", "doc-2"]


def test_semantic_selection_budget_runs_dry():
    from app.brain.extraction import select_semantic_chunks

    rows = [ {"source_id": "d1", "source_kind": "doc", "path": f"d{i}", "text": "x"} for i in range(3)]
    assert len(select_semantic_chunks(rows, 99)) == 3


async def test_pasted_text_is_user_provenance(db_session):
    """the user's own words cite provenance type user, not doc."""
    from app.service import add_text_source

    from app.models import Project, ProjectMember, User

    user = User(auth_id="t", email="t@t.dev")
    db_session.add(user)
    await db_session.commit()
    project = Project(name="text project")
    db_session.add(project)
    await db_session.flush()
    db_session.add(ProjectMember(project_id=project.id, user_id=user.id))
    await db_session.commit()

    src = await add_text_source(db_session, project.id, "we sell to jewelers", "my note")
    chunks = (await db_session.execute(
        __import__("sqlalchemy").select(__import__("app.models", fromlist=["SourceChunk"]).SourceChunk).where(
            __import__("app.models", fromlist=["SourceChunk"]).SourceChunk.source_id == src.id
        )
    )).scalars().all()
    assert chunks and all(c.kind == "user" for c in chunks)


async def test_doc_source_extraction_keeps_doc_provenance(db_session):
    from app.service import add_doc_source, create_project, extract_facts
    from conftest import FakeRun

    from app.models import Project, ProjectMember, User

    user = User(auth_id="t", email="t@t.dev")
    db_session.add(user)
    await db_session.commit()
    project = Project(name="docs project")
    db_session.add(project)
    await db_session.flush()
    db_session.add(ProjectMember(project_id=project.id, user_id=user.id))
    await db_session.commit()

    await add_doc_source(db_session, project.id, "readme.md", text="# growth\nthe readme names made-to-order jewelers as the first segment")
    canned = FakeRun({"extract_facts": {"facts": [{"kind": "audience", "statement": "the readme names made-to-order jewelers as the first segment", "lines": [1, 2]}]}})
    facts = await extract_facts(db_session, project.id, canned)
    semantic = [f for f in facts if f.kind == "audience"]
    assert semantic and semantic[0].provenance[0]["type"] == "doc"
    assert semantic[0].provenance[0]["ref"] == "readme.md"
